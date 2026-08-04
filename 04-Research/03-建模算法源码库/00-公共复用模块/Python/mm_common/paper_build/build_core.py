"""国赛论文 docx 生成核心（参数化版）。

从 06-Projects/2025国赛A题-烟幕干扰弹投放策略/Code/src/build_submission.py
提取并参数化（原文件未改动）。排版常量由 constants_2026 提供，
以便其他年份/赛区按 SPEC 调整；默认等价于 2025 项目的 Stage G 管线。

电子版口径：A4、页边距≥2.5cm、摘要第一页、页码从第 1 页起（单节+PAGE 域）、
无目录、无承诺书/编号页。公式以 matplotlib 渲染图片嵌入（T 级技术约束）。
"""

from __future__ import annotations

import hashlib
import re
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image

from . import constants_2026 as default_spec


def set_run_font(run, spec, cn=None, latin=None, size=None, bold=None, italic=None, color=None):
    cn = cn if cn is not None else spec.FONT_CN
    latin = latin if latin is not None else spec.FONT_LATIN
    size = size if size is not None else spec.FONT_SIZE_PT
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), cn)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, spec):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, val in zip(("top", "start", "bottom", "end"), spec.CELL_MARGIN_DXA):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    tr_pr.append(cant)


def add_page_number(paragraph, spec):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])
    set_run_font(run, spec, size=spec.PAGE_NUMBER_FONT_SIZE_PT)


def configure_document(doc: Document, spec, margin_cm: float | None = None):
    section = doc.sections[0]
    section.page_width = Cm(spec.PAGE_WIDTH_CM)
    section.page_height = Cm(spec.PAGE_HEIGHT_CM)
    margin = spec.MARGIN_CM if margin_cm is None else margin_cm
    section.top_margin = Cm(margin)
    section.bottom_margin = Cm(margin)
    section.left_margin = Cm(margin)
    section.right_margin = Cm(margin)
    section.header_distance = Cm(spec.HEADER_FOOTER_DISTANCE_CM)
    section.footer_distance = Cm(spec.HEADER_FOOTER_DISTANCE_CM)
    section.different_first_page_header_footer = False
    add_page_number(section.footer.paragraphs[0], spec)

    normal = doc.styles["Normal"]
    normal.font.name = spec.FONT_LATIN
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), spec.FONT_CN)
    normal.font.size = Pt(spec.FONT_SIZE_PT)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(spec.FIRST_LINE_INDENT_CM)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = spec.LINE_SPACING
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for level, size, before, after in spec.HEADING_LEVELS:
        name = f"Heading {level}"
        style = doc.styles[name]
        style.font.name = spec.FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), spec.HEADING_FONT_CN)
        style.font.size = Pt(size)
        style.font.bold = True
        pf = style.paragraph_format
        pf.keep_with_next = True
        pf.keep_together = True
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.first_line_indent = Cm(0)

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = spec.FONT_LATIN
        style._element.rPr.rFonts.set(qn("w:eastAsia"), spec.FONT_CN)
        style.font.size = Pt(spec.FONT_SIZE_PT)
        style.paragraph_format.left_indent = Cm(spec.FIRST_LINE_INDENT_CM)
        style.paragraph_format.first_line_indent = Cm(-0.37)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.15

    if "Caption CN" not in [s.name for s in doc.styles]:
        cap = doc.styles.add_style("Caption CN", WD_STYLE_TYPE.PARAGRAPH)
        cap.font.name = spec.FONT_LATIN
        cap._element.rPr.rFonts.set(qn("w:eastAsia"), spec.FONT_CN)
        cap.font.size = Pt(spec.CAPTION_FONT_SIZE_PT)
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.keep_with_next = True
        cap.paragraph_format.space_before = Pt(2)
        cap.paragraph_format.space_after = Pt(5)
        cap.paragraph_format.first_line_indent = Cm(0)

    if "Equation CN" not in [s.name for s in doc.styles]:
        eq = doc.styles.add_style("Equation CN", WD_STYLE_TYPE.PARAGRAPH)
        eq.font.name = "Cambria Math"
        eq._element.rPr.rFonts.set(qn("w:eastAsia"), "Cambria Math")
        eq.font.size = Pt(spec.FONT_SIZE_PT)
        eq.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        eq.paragraph_format.keep_together = True
        eq.paragraph_format.space_before = Pt(3)
        eq.paragraph_format.space_after = Pt(3)
        eq.paragraph_format.first_line_indent = Cm(0)


INLINE_PATTERN = re.compile(r"(\*\*.*?\*\*|`.*?`|\$.*?\$|\[[^\]]+\]\([^)]+\))")


def latex_inline_to_text(source: str, spec) -> str:
    """Convert compact inline LaTeX to readable text on the Word baseline."""
    text = source.strip()
    text = re.sub(r"\\rm\s+([A-Za-z]+)", r"\1", text)
    text = re.sub(r"\\(?:boldsymbol|mathbf|mathrm|mathcal)\s+", "", text)
    text = re.sub(
        r"\\(?:boldsymbol|mathbf|mathrm|mathcal)\{([^{}]*)\}",
        r"\1",
        text,
    )
    text = text.replace(r"\widehat\lambda", "λ̂")
    text = text.replace(r"\bar J", "J̄")
    for old, new in spec.INLINE_LATEX_MAP.items():
        text = text.replace(old, new)

    superscript_map = str.maketrans("0123456789+-", "⁰¹²³⁴⁵⁶⁷⁸⁹⁺⁻")

    def superscript(match):
        value = match.group(1)
        if all(ch in "0123456789+-" for ch in value):
            return value.translate(superscript_map)
        return f"^{value}"

    text = re.sub(r"\^\{([^{}]+)\}", superscript, text)
    text = re.sub(r"_\{([^{}]+)\}", r"_\1", text)
    text = text.replace("^*", "*")
    text = text.replace(r"\ ", " ")
    # Unknown commands remain readable and are caught by the residual-TeX QA.
    return re.sub(r"\\([A-Za-z]+)", r"\1", text)


def add_inline(paragraph, text: str, spec, base_size=None):
    base_size = base_size if base_size is not None else spec.FONT_SIZE_PT
    pos = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > pos:
            set_run_font(paragraph.add_run(text[pos : match.start()]), spec, size=base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, spec, size=base_size, bold=True)
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, spec, cn="等线", latin="Consolas", size=max(base_size - 1, 8.5))
        elif token.startswith("$"):
            run = paragraph.add_run(latex_inline_to_text(token[1:-1], spec))
            set_run_font(run, spec, cn="Cambria Math", latin="Cambria Math", size=base_size)
        else:
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token).groups()
            run = paragraph.add_run(f"{label}（{url}）")
            set_run_font(run, spec, size=base_size)
            run.font.color.rgb = RGBColor(31, 78, 121)
        pos = match.end()
    if pos < len(text):
        set_run_font(paragraph.add_run(text[pos:]), spec, size=base_size)


def clean_markdown_text(text: str) -> str:
    text = re.sub(r"\[\[([^\]|]+)\|([^\]]+)\]\]", r"\2", text)
    text = re.sub(r"\[\[([^\]]+)\]\]", r"\1", text)
    return text.replace("\\_", "_")


def add_table(doc: Document, rows: list[list[str]], spec):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    table.alignment = 1
    table.autofit = True
    for r_idx, values in enumerate(rows):
        row = table.rows[r_idx]
        set_cant_split(row)
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx in range(cols):
            cell = row.cells[c_idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell, spec)
            if r_idx == 0:
                set_cell_shading(cell, spec.TABLE_HEADER_FILL)
            para = cell.paragraphs[0]
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER if r_idx == 0 else (
                WD_ALIGN_PARAGRAPH.LEFT if c_idx == 0 else WD_ALIGN_PARAGRAPH.CENTER
            )
            para.paragraph_format.first_line_indent = Cm(0)
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.0
            val = values[c_idx] if c_idx < len(values) else ""
            add_inline(para, val, spec, base_size=spec.TABLE_FONT_SIZE_PT)
            if r_idx == 0:
                for run in para.runs:
                    run.bold = True
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(0)
    after.paragraph_format.line_spacing = 0.5


def normalized_image_path(image_path: Path, spec, temp_label: str) -> Path:
    """Flatten alpha and cap raster dimensions for stable Word 2016 import."""
    digest = hashlib.sha256(
        f"{image_path.resolve()}:{image_path.stat().st_mtime_ns}".encode("utf-8")
    ).hexdigest()[:16]
    target_dir = Path(tempfile.gettempdir()) / temp_label
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"{image_path.stem}-{digest}.jpg"
    if not target.exists():
        with Image.open(image_path) as image:
            image = image.convert("RGB")
            image.thumbnail(spec.IMAGE_MAX_PX, Image.Resampling.LANCZOS)
            image.save(
                target,
                format="JPEG",
                quality=spec.IMAGE_QUALITY,
                optimize=True,
                dpi=(spec.IMAGE_DPI, spec.IMAGE_DPI),
            )
    return target


def equation_image_path(latex: str, spec, temp_label: str) -> Path:
    """Render display math as a high-resolution image when OMML conversion is unavailable."""
    source = " ".join(part.strip() for part in latex.splitlines()).strip()
    # Preserve a token boundary when LaTeX uses ``\\boldsymbol P`` without
    # braces.  Removing the command naively turns ``\\forall\\boldsymbol P``
    # into the invalid MathText token ``\\forallP``.
    source = re.sub(r"\\boldsymbol\s*\{([^{}]+)\}", r"{\1}", source)
    source = re.sub(r"\\boldsymbol\s+([A-Za-z])", r"{\1}", source)
    source = re.sub(r"\\boldsymbol\s*", "", source)
    source = source.replace(r"\frac12", r"\frac{1}{2}")
    source = re.sub(r"\\le(?![A-Za-z])", r"\\leq", source)
    source = re.sub(r"\\ge(?![A-Za-z])", r"\\geq", source)
    source = source.replace(r"\mathsf T", r"\mathsf{T}")
    source = source.replace(r"\bar J", r"\bar{J}")
    source = source.replace(r"\sqrt n", r"\sqrt{n}")
    source = source.replace(r"\mathbb E", r"\mathbb{E}")
    source = re.sub(r"\\mathcal\s+([A-Za-z])", r"\\mathcal{\1}", source)
    source = re.sub(r"\\text\{([^{}]+)\}", r"\\mathrm{\1}", source)
    source = re.sub(r"\\operatorname\{([^{}]+)\}", r"\\mathrm{\1}", source)
    digest = hashlib.sha256(source.encode("utf-8")).hexdigest()[:16]
    target_dir = Path(tempfile.gettempdir()) / temp_label
    target_dir.mkdir(parents=True, exist_ok=True)
    target = target_dir / f"equation-{digest}.png"
    if not target.exists():
        import matplotlib

        matplotlib.use("Agg")
        import matplotlib.pyplot as plt

        fig = plt.figure(figsize=(8, 0.55), dpi=300, facecolor="white")
        fig.text(
            0.5,
            0.5,
            f"${source}$",
            ha="center",
            va="center",
            fontsize=12,
            color="black",
        )
        fig.savefig(
            target,
            dpi=spec.IMAGE_DPI,
            bbox_inches="tight",
            pad_inches=0.025,
            facecolor="white",
        )
        plt.close(fig)
    return target


def add_equation(doc: Document, latex: str, number: int, spec, temp_label: str):
    p = doc.add_paragraph(style="Equation CN")
    p.paragraph_format.keep_together = True
    p.paragraph_format.tab_stops.add_tab_stop(Cm(spec.EQUATION_TAB_CM), WD_TAB_ALIGNMENT.RIGHT)
    try:
        image_path = equation_image_path(latex, spec, temp_label)
        with Image.open(image_path) as image:
            native_width_cm = image.width / spec.IMAGE_DPI * 2.54
        width = Cm(min(spec.EQUATION_WIDTH_MAX_CM, max(spec.EQUATION_WIDTH_MIN_CM, native_width_cm)))
        p.add_run().add_picture(str(image_path), width=width)
    except Exception:
        # Last-resort readable fallback; never silently drop a formula.
        fallback = " ".join(part.strip() for part in latex.splitlines())
        run = p.add_run(fallback)
        set_run_font(run, spec, cn="Cambria Math", latin="Cambria Math", size=spec.EQUATION_FONT_SIZE_PT)
    p.add_run("\t")
    number_run = p.add_run(f"({number})")
    set_run_font(number_run, spec, size=spec.EQUATION_FONT_SIZE_PT)


def add_figure(doc: Document, image_path: Path, caption: str, figure_number: int, spec, temp_label: str):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run()
    run.add_picture(str(normalized_image_path(image_path, spec, temp_label)), width=Cm(spec.FIGURE_WIDTH_CM))
    cap = doc.add_paragraph(style="Caption CN")
    add_inline(cap, f"图 {figure_number}  {caption}", spec, base_size=spec.CAPTION_FONT_SIZE_PT)


def parse_table(lines: list[str], start: int):
    rows = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        raw = [c.strip() for c in lines[i].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", c or "") for c in raw):
            rows.append(raw)
        i += 1
    return rows, i


def build_from_markdown(
    md_path: Path,
    out_path: Path,
    *,
    spec=None,
    ai_detail=False,
    include_images=True,
    image_limit: int | None = None,
    margin_cm: float | None = None,
    title: str | None = None,
    subject: str | None = None,
    temp_label: str = "mm-common-docx-images",
):
    """Render an UTF-8 Markdown paper to a formatted .docx.

    md_path: 论文 Markdown（相对图片路径以 md 所在目录解析）。
    out_path: 输出 docx 路径。
    spec: 排版常量模块（默认 constants_2026）。
    ai_detail: AI 使用详情文档模式（跳过未填写的记录模板）。
    include_images / image_limit: 图片嵌入开关与上限。
    margin_cm: 覆盖页边距（cm）。
    title / subject: docx core properties 覆盖（默认为匿名+赛题通用描述）。
    temp_label: 图片归一化缓存目录名。
    """
    spec = spec if spec is not None else default_spec
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc, spec, margin_cm=margin_cm)

    # Remove YAML front matter when present.
    if lines and lines[0].strip() == "---":
        try:
            end = lines.index("---", 1)
            lines = lines[end + 1 :]
        except ValueError:
            pass

    in_equation = False
    equation_lines: list[str] = []
    in_code = False
    code_lines: list[str] = []
    figure_no = 0
    equation_no = 0
    first_title = True
    skip_ai_template = False
    i = 0
    while i < len(lines):
        raw = lines[i]
        line = clean_markdown_text(raw.strip())
        if ai_detail and line.startswith("### 记录 NNN"):
            skip_ai_template = True
            i += 1
            continue
        if skip_ai_template:
            if line.startswith("## 正文与参考文献待办"):
                skip_ai_template = False
            else:
                i += 1
                continue
        if line == "<!-- PAGEBREAK -->":
            doc.add_page_break()
            i += 1
            continue
        if line.startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.left_indent = Cm(0.6)
                p.paragraph_format.right_indent = Cm(0.4)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(3)
                for idx, code_line in enumerate(code_lines):
                    if idx:
                        p.add_run().add_break()
                    run = p.add_run(code_line)
                    set_run_font(run, spec, cn="等线", latin="Consolas", size=8)
                in_code = False
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if line == "$$":
            if not in_equation:
                in_equation = True
                equation_lines = []
            else:
                equation_no += 1
                add_equation(doc, "\n".join(equation_lines), equation_no, spec, temp_label)
                in_equation = False
            i += 1
            continue
        if in_equation:
            equation_lines.append(raw)
            i += 1
            continue
        if not line:
            i += 1
            continue
        if line.startswith("> [!"):
            i += 1
            continue
        if line.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.right_indent = Cm(0.4)
            p.paragraph_format.first_line_indent = Cm(0)
            add_inline(p, line.lstrip("> ").strip(), spec)
            i += 1
            continue
        if line.startswith("# "):
            title_text = line[2:].strip()
            if first_title:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Cm(0)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after = Pt(spec.TITLE_SPACE_AFTER_PT)
                p.paragraph_format.keep_with_next = True
                run = p.add_run(title_text)
                set_run_font(run, spec, cn=spec.HEADING_FONT_CN, size=spec.TITLE_FONT_SIZE_PT, bold=True)
                first_title = False
            else:
                doc.add_heading(title_text, level=1)
            i += 1
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=1)
            i += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=2)
            i += 1
            continue
        if line.startswith("#### "):
            doc.add_heading(line[5:].strip(), level=3)
            i += 1
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", line)
        if image_match:
            figure_no += 1
            caption, rel = image_match.groups()
            image_path = (md_path.parent / rel).resolve()
            if include_images and (image_limit is None or figure_no <= image_limit):
                add_figure(doc, image_path, caption, figure_no, spec, temp_label)
            else:
                p = doc.add_paragraph(style="Caption CN")
                add_inline(p, f"图 {figure_no}  {caption}（调试构建未嵌图）", spec, base_size=spec.CAPTION_FONT_SIZE_PT)
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows, spec)
            continue
        num_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if num_match:
            # Preserve the source number explicitly. Word otherwise continues
            # one automatic list across unrelated sections (for example 7–10
            # instead of restarting a later four-step algorithm at 1).
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(spec.FIRST_LINE_INDENT_CM)
            p.paragraph_format.first_line_indent = Cm(-0.37)
            p.paragraph_format.line_spacing = 1.15
            add_inline(p, f"{num_match.group(1)}. {num_match.group(2)}", spec)
            i += 1
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, line[2:], spec)
            i += 1
            continue

        p = doc.add_paragraph()
        if line.startswith("**摘要：**") or line.startswith("**关键词：**") or line.startswith("**AI 辅助声明：**"):
            p.paragraph_format.first_line_indent = Cm(0)
        is_ai_note = line.startswith("AI辅助标注：") or line.startswith("AI 辅助标注：")
        if is_ai_note:
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(0.5)
            p.paragraph_format.right_indent = Cm(0.5)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
        add_inline(p, line, spec)
        if is_ai_note:
            for run in p.runs:
                run.font.color.rgb = RGBColor(*spec.AI_ANNOTATION_RGB)
                run.italic = True
                run.font.size = Pt(spec.AI_ANNOTATION_SIZE_PT)
        i += 1

    # Metadata must be anonymous for competition submission artifacts.
    props = doc.core_properties
    props.title = title if title is not None else ("AI 工具使用详情" if ai_detail else "数学建模竞赛论文")
    props.author = ""
    props.last_modified_by = ""
    props.subject = subject if subject is not None else "全国大学生数学建模竞赛论文"
    props.keywords = ""
    props.comments = ""
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    return out_path
