"""国赛论文 docx 机器 QA（14 项）。

逐项输出 PASS/FAIL + 证据；--report-json 输出 JSON。任一项 FAIL 时
exit code 非 0（供管线阻断）。页数/版面等需渲染的项为人工项，
本模块不做机器判定（见审计清单 A5 与 skill 已知限制）。

检查项与 04-Research/08-论文写作与复现/论文交付审计清单.md 对应：
A1 摘要第一页、A3 页码域、A4 无目录、A6 纸张边距、A7 正文字体、
A8 公式编号、A10 图注连续性+图嵌入、A11 文献双向对应、A13 文件大小、
A15 匿名 core、C10 表头重复、C1 残余 LaTeX。
"""

from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

from . import constants_2026 as spec


def _all_text(doc: Document) -> str:
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def check_page_size(doc: Document, spec) -> tuple[bool, str]:
    section = doc.sections[0]
    w_ok = abs(section.page_width.cm - spec.PAGE_WIDTH_CM) < 0.01
    h_ok = abs(section.page_height.cm - spec.PAGE_HEIGHT_CM) < 0.01
    return (w_ok and h_ok), f"页面 {section.page_width.cm:.2f}×{section.page_height.cm:.2f} cm（要求 {spec.PAGE_WIDTH_CM}×{spec.PAGE_HEIGHT_CM}）"


def check_margins(doc: Document, spec) -> tuple[bool, str]:
    section = doc.sections[0]
    margins = (section.top_margin.cm, section.bottom_margin.cm, section.left_margin.cm, section.right_margin.cm)
    ok = all(m >= spec.MARGIN_CM - 0.01 for m in margins)
    return ok, f"边距 {margins} cm（要求 ≥{spec.MARGIN_CM}）"


def check_page_number_field(doc: Document, spec) -> tuple[bool, str]:
    xml = doc.sections[0].footer.paragraphs[0]._element.xml
    return "PAGE" in xml, "页脚包含 PAGE 域" if "PAGE" in xml else "页脚缺少 PAGE 域"


def check_no_toc(doc: Document, spec) -> tuple[bool, str]:
    toc_fields = re.findall(r'TOC[^<]*', doc.element.xml)
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    has_toc_heading = any(h.strip() == "目录" for h in headings)
    return (not toc_fields and not has_toc_heading), (
        f"无 TOC 域" if not toc_fields else f"发现 TOC 域 {len(toc_fields)} 处"
    ) + ("；无目录标题" if not has_toc_heading else "；存在目录标题")


def check_residual_latex(doc: Document, spec) -> tuple[bool, str]:
    text = _all_text(doc)
    hits = re.findall(r"\\(?:frac|boldsymbol|mathbf|mathrm|mathbb|mathcal|forall|exists|sum|leq|geq|lambda|theta|tau)[A-Za-z]*", text)
    return not hits, "无残余 LaTeX 命令" if not hits else f"残余 LaTeX 命令 {len(hits)} 处：{hits[:5]}"


def check_caption_sequence(doc: Document, spec) -> tuple[bool, str]:
    numbers = []
    for p in doc.paragraphs:
        # 只统计 Caption CN 样式（真实图题）；正文"图 N 中…"开头的句子不算图题
        if p.style.name != "Caption CN":
            continue
        m = re.fullmatch(r"图\s+(\d+)\s+.*", p.text)
        if m:
            numbers.append(int(m.group(1)))
    expected = list(range(1, len(numbers) + 1))
    return numbers == expected, f"图题编号连续（图 {len(numbers)} 张，编号 {numbers[:3]}…）" if numbers == expected else f"图题编号跳号：{numbers}"


def check_equation_numbering(doc: Document, spec) -> tuple[bool, str]:
    numbers = []
    for p in doc.paragraphs:
        m = re.search(r"\((\d+)\)\s*$", p.text.strip())
        if m and p.style.name == "Equation CN":
            numbers.append(int(m.group(1)))
    expected = list(range(1, len(numbers) + 1))
    return numbers == expected, f"公式编号连续（{len(numbers)} 个）" if numbers == expected else f"公式编号跳号：{numbers}"


def check_image_count(doc: Document, spec, expected: int | None = None) -> tuple[bool, str]:
    count = len(doc.inline_shapes)
    if expected is not None:
        ok = count == expected
        return ok, f"嵌入图片 {count} 张（期望 {expected}）"
    return count >= 1, f"嵌入图片 {count} 张"


def check_anonymous_core(doc: Document, spec) -> tuple[bool, str]:
    props = doc.core_properties
    ok = not props.author and not props.last_modified_by
    return ok, f"author='{props.author}' last_modified_by='{props.last_modified_by}'（要求均为空）"


def check_file_size(path: Path, spec) -> tuple[bool, str]:
    mb = path.stat().st_size / (1024 * 1024)
    return mb <= spec.MAX_FILE_SIZE_MB, f"文件 {mb:.2f} MB（要求 ≤{spec.MAX_FILE_SIZE_MB}）"


def check_body_font(doc: Document, spec) -> tuple[bool, str]:
    normal = doc.styles["Normal"]
    east_asia = normal._element.rPr.rFonts.get(qn("w:eastAsia")) if normal._element.rPr is not None and normal._element.rPr.rFonts is not None else None
    size = normal.font.size.pt if normal.font.size else None
    ok = east_asia == spec.FONT_CN and size == spec.FONT_SIZE_PT
    return ok, f"Normal 样式 eastAsia={east_asia} 字号={size}pt（要求 {spec.FONT_CN} {spec.FONT_SIZE_PT}pt）"


def check_table_header_repeat(doc: Document, spec) -> tuple[bool, str]:
    if not doc.tables:
        return True, "无表格（无需跨页重复表头，人工确认）"
    first = doc.tables[0]
    first_row = first.rows[0]
    has_repeat = "tblHeader" in first_row._tr.xml
    return has_repeat, "表 1 首行设置跨页重复表头" if has_repeat else "表 1 首行未设置跨页重复表头"


def check_reference_bidirection(doc: Document, spec) -> tuple[bool, str]:
    text = _all_text(doc)
    # 文献列表 = "参考文献"标题到下一个章节标题（如"附录"）之间的块；
    # 其余全部算正文引用区。否则条目自身的 [n] 会被误算为正文引用，
    # 且置于参考文献之后的附录内引用也会被误判为未引用。
    ref_marker = re.search(r"^参考文献[^\n]*$", text, re.MULTILINE)
    if ref_marker is None:
        return True, "未找到'参考文献'节（人工确认文档结构）"
    tail = text[ref_marker.start() :]
    next_section = re.search(r"^(?:#|附录)", tail, re.MULTILINE)
    refs_text = tail if next_section is None else tail[: next_section.start()]
    body = text[: ref_marker.start()] + (tail if next_section is None else tail[next_section.start() :])
    cited = set(int(m.group(1)) for m in re.finditer(r"\[(\d+)\]", body))
    refs = set(
        int(m.group(1))
        for line in refs_text.splitlines()
        if (m := re.match(r"^\s*\[(\d+)\]\s+\S", line.strip()))
    )
    if not refs:
        return True, "参考文献节无条目（人工确认）"
    missing_refs = refs - cited
    dangling_cites = cited - refs
    if missing_refs or dangling_cites:
        return False, f"文献 {sorted(refs)}；正文引用 {sorted(cited)}；未引用条目 {sorted(missing_refs)}；无条目引用 {sorted(dangling_cites)}"
    return True, f"文献 {len(refs)} 条全部被正文引用，无悬挂引用"


def check_summary_first_page(doc: Document, spec) -> tuple[bool, str]:
    first_page_text = ""
    for p in doc.paragraphs[:12]:
        first_page_text += p.text
    return "摘要" in first_page_text, "前 12 段含'摘要'" if "摘要" in first_page_text else "前 12 段未见'摘要'（检查标题与摘要是否同页）"


CHECKS = [
    ("page-size", check_page_size),
    ("margins", check_margins),
    ("page-number-field", check_page_number_field),
    ("no-toc", check_no_toc),
    ("residual-latex", check_residual_latex),
    ("caption-sequence", check_caption_sequence),
    ("equation-numbering", check_equation_numbering),
    ("image-count", check_image_count),
    ("anonymous-core", check_anonymous_core),
    ("file-size", check_file_size),
    ("body-font", check_body_font),
    ("table-header-repeat", check_table_header_repeat),
    ("reference-bidirection", check_reference_bidirection),
    ("summary-first-page", check_summary_first_page),
]


def run_qa(docx_path: Path, expected_images: int | None = None) -> list[dict]:
    doc = Document(str(docx_path))
    results = []
    for name, fn in CHECKS:
        try:
            if name == "file-size":
                ok, evidence = fn(docx_path, spec)
            elif name == "image-count":
                ok, evidence = fn(doc, spec, expected_images)
            else:
                ok, evidence = fn(doc, spec)
        except Exception as exc:  # noqa: BLE001 - QA 单项失败记录为 FAIL 而非中断
            ok, evidence = False, f"检查异常：{exc!r}"
        results.append({"check": name, "status": "PASS" if ok else "FAIL", "evidence": evidence})
    return results


def main():
    parser = argparse.ArgumentParser(description="国赛论文 docx 机器 QA（14 项）")
    parser.add_argument("docx", type=Path, help="待检查的 docx 路径")
    parser.add_argument("--expected-images", type=int, default=None, help="期望嵌入图片数")
    parser.add_argument("--report-json", type=Path, default=None, help="输出 JSON 报告路径")
    parser.add_argument("--ai-detail", action="store_true", help="AI 使用详情文档模式：跳过摘要第一页检查")
    args = parser.parse_args()

    results = run_qa(args.docx, expected_images=args.expected_images)
    if args.ai_detail:
        for r in results:
            if r["check"] == "summary-first-page":
                r["status"] = "PASS"
                r["evidence"] = "AI 详情文档无摘要第一页要求，跳过"
    failures = [r for r in results if r["status"] != "PASS"]
    for r in results:
        print(f"[{r['status']}] {r['check']}: {r['evidence']}")
    print(f"结论：{len(results) - len(failures)}/{len(results)} PASS")
    if failures:
        print("阻断：存在 FAIL 项，修复后重跑")
    if args.report_json is not None:
        payload = {
            "docx": str(args.docx.resolve()),
            "results": results,
            "pass": len(results) - len(failures),
            "total": len(results),
            "fail": [r["check"] for r in failures],
            "manual_items": ["page-count(≤30页)", "layout-render", "actual-tool-version"],
        }
        args.report_json.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"报告：{args.report_json}")
    return 1 if failures else 0


if __name__ == "__main__":
    raise SystemExit(main())
