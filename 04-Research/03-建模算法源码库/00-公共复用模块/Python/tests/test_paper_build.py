"""paper_build 管线单元测试（unittest，与 mm_common/tests 风格一致）。

覆盖：行内 LaTeX 转换、公式编号、图题编号、边距、匿名 core、
页码域、残余 LaTeX、机器 QA 全项。
"""

from __future__ import annotations

import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from mm_common.paper_build import constants_2026 as spec  # noqa: E402
from mm_common.paper_build.build_core import (  # noqa: E402
    build_from_markdown,
    latex_inline_to_text,
)
from mm_common.paper_build.paper_qa import run_qa  # noqa: E402

FIXTURE = """---
title: 测试论文
---

# 测试论文标题

**摘要：**针对测试问题，本文建立模型，得到结果 1.234 s，并通过测试验证。

## 1 问题

设 $\\lambda$ 为参数，则 $\\boldsymbol P$ 满足 $\\forall k\\in\\mathcal K$。该处理见 [1]。

$$
J=\\sum_{k} d_k \\le 10
$$

如图 1 所示。

![测试图](missing.png)

$$
\\mu(E)=\\int_0^1 dt
$$

## 2 参考文献

[1] 测试文献. 测试期刊[J]. 2026.
"""


class LatexInlineTest(unittest.TestCase):
    def test_unicode_map(self):
        # 与 build_core.add_inline 一致：调用前已剥去 $ 包裹符
        self.assertEqual(latex_inline_to_text(r"\lambda \times \mu", spec), "λ × μ")

    def test_boldsymbol_removed(self):
        self.assertEqual(latex_inline_to_text(r"\boldsymbol P", spec), "P")

    def test_residual_commands_removed(self):
        # 与 build_submission.py 一致：命令替换后保留原文中的空格
        self.assertEqual(latex_inline_to_text(r"\forall k\in\mathcal K", spec), "∀ k∈K")


class BuildTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        (self.root / "figures").mkdir()
        md = self.root / "paper.md"
        md.write_text(FIXTURE, encoding="utf-8")
        self.md = md
        self.out = self.root / "paper.docx"

    def tearDown(self):
        self.tmp.cleanup()

    def build(self, **kwargs):
        return build_from_markdown(
            self.md, self.out, spec=spec, include_images=False, **kwargs
        )

    def test_build_produces_docx(self):
        self.build()
        self.assertTrue(self.out.exists())

    def test_equation_numbering_sequential(self):
        self.build()
        doc = Document(str(self.out))
        numbers = [
            int(p.text.strip().rsplit("(", 1)[1][:-1])
            for p in doc.paragraphs
            if p.style.name == "Equation CN" and p.text.strip().endswith(")")
        ]
        self.assertEqual(numbers, [1, 2])

    def test_caption_sequence(self):
        self.build()
        doc = Document(str(self.out))
        captions = [p.text for p in doc.paragraphs if p.text.startswith("图 ")]
        self.assertEqual(captions, ["图 1  测试图（调试构建未嵌图）"])

    def test_margins_at_least_2_5cm(self):
        self.build()
        section = Document(str(self.out)).sections[0]
        for m in (section.top_margin.cm, section.bottom_margin.cm, section.left_margin.cm, section.right_margin.cm):
            self.assertGreaterEqual(m, 2.49)

    def test_anonymous_core(self):
        self.build()
        props = Document(str(self.out)).core_properties
        self.assertEqual(props.author, "")
        self.assertEqual(props.last_modified_by, "")

    def test_page_number_field(self):
        self.build()
        xml = Document(str(self.out)).sections[0].footer.paragraphs[0]._element.xml
        self.assertIn("PAGE", xml)

    def test_no_residual_latex_in_docx(self):
        self.build()
        text = "\n".join(p.text for p in Document(str(self.out)).paragraphs)
        self.assertNotIn("\\lambda", text)
        self.assertNotIn("\\forall", text)

    def test_ai_note_gray_italic(self):
        md = self.root / "ai.md"
        md.write_text("# AI 详情\n\nAI辅助标注：本段为 AI 生成。", encoding="utf-8")
        out = self.root / "ai.docx"
        build_from_markdown(md, out, spec=spec, ai_detail=True, include_images=False)
        doc = Document(str(out))
        for p in doc.paragraphs:
            if p.text.startswith("AI辅助标注"):
                for run in p.runs:
                    self.assertTrue(run.italic)
                break
        else:
            self.fail("未找到 AI 标注段落")


class QaTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        (self.root / "figures").mkdir()
        md = self.root / "paper.md"
        md.write_text(FIXTURE, encoding="utf-8")
        self.out = self.root / "paper.docx"
        build_from_markdown(md, self.out, spec=spec, include_images=False)

    def tearDown(self):
        self.tmp.cleanup()

    def test_qa_all_pass(self):
        # 调试构建不含图，但两条 $$ 公式仍渲染为 matplotlib 图片 → 期望 2 张
        results = run_qa(self.out, expected_images=2)
        failed = [r for r in results if r["status"] != "PASS"]
        self.assertEqual(failed, [], f"FAIL 项：{[r['evidence'] for r in failed]}")


if __name__ == "__main__":
    unittest.main()
